"""Real local commands can use their workspace without widening protected reads."""
import json
import os
from pathlib import Path
import shlex
import shutil
import subprocess
import sys

import pytest

from athena.auth import Auth
from athena.config import Config
from athena.protection import command_environment, protected_command


@pytest.fixture
def config(tmp_path):
    config = Config(data_dir=tmp_path / "Application Support" / "Athena Test",
                    repo_dir=tmp_path / "Core App" / "Resources")
    config.prepare()
    config.repo_dir.mkdir(parents=True)
    return config


def test_command_environment_preserves_venv_path_and_filters_operator_values(tmp_path, monkeypatch):
    executable = tmp_path / "venv with spaces" / "bin" / "python"
    executable.parent.mkdir(parents=True)
    executable.symlink_to(sys.executable)
    monkeypatch.setattr(sys, "executable", str(executable))
    inherited = {"DEEPSEEK_API_KEY": "test-provider-secret", "ATHENA_AUDIO_API_KEY": "test-voice-secret",
                 "PYTHONPATH": "/injected/imports", "PYTHONHOME": "/injected/home",
                 "PYTHONSTARTUP": "/injected/startup.py", "ATHENA_PYTHON": "/injected/python"}
    for name, value in inherited.items():
        monkeypatch.setenv(name, value)
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    env = command_environment(tmp_path / "files")
    assert env["PATH"] == str(executable.parent) + os.pathsep + "/usr/bin:/bin"
    assert env["ATHENA_PYTHON"] == str(executable)  # Do not resolve out of the venv.
    assert env["ATHENA_WORKSPACE"] == str(tmp_path / "files")
    assert env["PYTHONNOUSERSITE"] == "1" and env["PYTHONDONTWRITEBYTECODE"] == "1"
    assert not any(value in json.dumps(env) for value in inherited.values())
    assert not {"PYTHONPATH", "PYTHONHOME", "PYTHONSTARTUP", "DEEPSEEK_API_KEY", "ATHENA_AUDIO_API_KEY"} & env.keys()


local_sandbox = pytest.mark.skipif(
    sys.platform != "darwin" or not shutil.which("sandbox-exec"), reason="macOS Seatbelt integration")


@local_sandbox
async def test_exec_cd_pwd_and_relative_excel_use_shipped_python(engine, store, config, monkeypatch):
    # A launch from Finder need not have uv or a project venv in its PATH.
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    monkeypatch.setenv("PYTHONPATH", "/must-not-be-inherited")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "must-not-be-inherited")
    user, _ = await Auth(store).register("workspace-user", "fixture-password")
    session = await store.session(user_id=user["id"])
    program = '''import json, os, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
assert "DEEPSEEK_API_KEY" not in os.environ and "PYTHONPATH" not in os.environ
book = Workbook()
sheet = book.active
sheet.append(["项目", "金额"])
sheet.append(["餐费", 28.5])
sheet.append(["交通", 12])
sheet.append(["合计", "=SUM(B2:B3)"])
book.save("日常开支.xlsx")
saved = load_workbook("日常开支.xlsx")
assert saved.active["B2"].value == 28.5 and saved.active["B4"].value == "=SUM(B2:B3)"
assert Path.cwd() == Path(os.environ["ATHENA_WORKSPACE"])
assert sys.executable == os.environ["ATHENA_PYTHON"] or Path(sys.executable).samefile(os.environ["ATHENA_PYTHON"])
assert not sys._is_gil_enabled()
print(json.dumps({"file": "日常开支.xlsx", "gil_enabled": sys._is_gil_enabled()}, ensure_ascii=False))
'''
    result = await engine.tools.shell({"command": 'cd "$ATHENA_WORKSPACE" && pwd -P && python3 -c ' + shlex.quote(program)}, session, None)
    assert result["exit_code"] == 0, result
    workspace = config.data_dir / "files" / user["id"] / session
    lines = result["output"].splitlines()
    assert Path(lines[0]) == workspace.resolve()
    assert json.loads(lines[-1]) == {"file": "日常开支.xlsx", "gil_enabled": False}
    assert (workspace / "日常开支.xlsx").is_file()
    # Explicit interpreter selection also works with spaces and no system deps.
    explicit = await engine.tools.shell({"command": '"$ATHENA_PYTHON" -c "import openpyxl; print(openpyxl.__version__)"'}, session, None)
    assert explicit["exit_code"] == 0, explicit
    assert explicit["output"].strip().startswith("3.")


@local_sandbox
async def test_new_excel_keeps_formula_and_verified_cache_in_user_preview(engine, store, config, monkeypatch):
    from athena.spreadsheet_preview import xlsx_preview
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    session = await store.session()
    code = '''import json, sys
from decimal import Decimal
import xlsxwriter
from openpyxl import load_workbook
amounts = [Decimal("12"), Decimal("35"), Decimal("20")]
with xlsxwriter.Workbook("带公式的开支.xlsx") as book:
    sheet = book.add_worksheet("开支")
    money = book.add_format({"num_format": "0.00"})
    sheet.write_row("A1", ["项目", "金额"])
    for row, (name, amount) in enumerate(zip(["早餐", "午餐", "交通"], amounts), 1):
        sheet.write(row, 0, name)
        sheet.write_number(row, 1, float(amount), money)
    sheet.write("A5", "合计")
    sheet.write_formula("B5", "=SUM(B2:B4)", money, float(sum(amounts)))
formula = load_workbook("带公式的开支.xlsx", data_only=False).active["B5"].value
cached = load_workbook("带公式的开支.xlsx", data_only=True).active["B5"].value
assert formula == "=SUM(B2:B4)" and cached == 67
assert not sys._is_gil_enabled()
print(json.dumps({"formula": formula, "cached": cached}))
'''
    result = await engine.tools.shell({"command": 'cd "$ATHENA_WORKSPACE" && python3 -c ' + shlex.quote(code)}, session, None)
    assert result["exit_code"] == 0, result
    assert json.loads(result["output"]) == {"formula": "=SUM(B2:B4)", "cached": 67}
    output = config.data_dir / "files" / "legacy" / session / "带公式的开支.xlsx"
    preview = xlsx_preview(output.read_bytes())
    sheet = preview["sheets"][0]
    cell = sheet["cells"][sheet["row_numbers"].index(5)][1]
    assert cell["kind"] == "formula" and cell["formula"].lstrip("=") == "SUM(B2:B4)"
    assert cell["formula_status"] == "saved" and float(cell["raw"]) == 67
    assert cell["display"] == "67.00"
    assert preview["formula_results"] == "not_recalculated"  # Preview must not claim to be an engine.


@local_sandbox
def test_workspace_traversal_keeps_data_contents_and_other_accounts_protected(config, tmp_path, monkeypatch):
    home = tmp_path / "Synthetic Home"
    (home / ".secrets").mkdir(parents=True)
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: home))
    area = config.data_dir / "files" / "alice"
    workspace = area / "session"
    workspace.mkdir(parents=True)
    extension = config.data_dir / "extensions" / "alice"
    extension.mkdir()
    owned_tool = extension / "tool.py"
    owned_tool.write_text("print('owned extension')")
    database = config.data_dir / "athena.sqlite3"
    database.write_text("synthetic-private-data")
    secret = home / ".secrets" / "provider.txt"
    secret.write_text("synthetic-private-data")
    envfile = config.repo_dir / ".env"
    envfile.write_text("synthetic-private-data")
    source = config.repo_dir / "source.py"
    source.write_text("original-source")
    foreign = config.data_dir / "files" / "bob" / "session" / "private.txt"
    foreign.parent.mkdir(parents=True)
    foreign.write_text("synthetic-private-data")
    foreign_extension = config.data_dir / "extensions" / "bob" / "tool.py"
    foreign_extension.parent.mkdir()
    foreign_extension.write_text("synthetic-private-data")
    alias = workspace / "database-link"
    alias.symlink_to(database)
    checks = [("list", path) for path in (config.data_dir, area.parent, extension.parent)]
    checks += [("read", path) for path in (database, secret, envfile, foreign, foreign_extension, alias)]
    checks += [("write", path) for path in (source, database, foreign, owned_tool)]
    code = '''import json, os, sys
from pathlib import Path
for operation, name in json.loads(sys.argv[1]):
    try:
        if operation == "list": os.listdir(name)
        elif operation == "read": Path(name).read_bytes()
        else: Path(name).write_text("changed")
    except PermissionError: pass
    else: raise AssertionError("protected operation was permitted: " + operation + " " + name)
print("all protected operations denied")
'''
    command = 'cd "$ATHENA_WORKSPACE" && python3 -c ' + shlex.quote(code) + " " + shlex.quote(json.dumps([(op, str(path)) for op, path in checks]))
    result = subprocess.run(protected_command(command, config, [area, extension], [area]),
                            cwd=workspace, env=command_environment(workspace), capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "all protected operations denied"
    assert source.read_text() == "original-source"
    assert owned_tool.read_text() == "print('owned extension')"
    allowed = subprocess.run(protected_command('cd ' + shlex.quote(str(extension)) + ' && "$ATHENA_PYTHON" tool.py',
                                               config, [area, extension], [area]),
                             cwd=workspace, env=command_environment(workspace), capture_output=True, text=True)
    assert allowed.returncode == 0 and allowed.stdout.strip() == "owned extension", allowed.stderr
