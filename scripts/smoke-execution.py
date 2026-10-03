"""Run inside a disposable web-like container beside the Linux execution broker."""
import asyncio
import json
import os
from pathlib import Path
import shlex
import sys
import re
import uuid

from athena.auth import Auth
from athena.config import Config, require_free_threading
from athena.files import discover_files, initialize_files, _open_file
from athena.store import Store
from athena.tools import ToolRunner


async def main():
    self_host_url = os.environ.get('ATHENA_TEST_SELF_HOST_URL')
    if not self_host_url:
        raise RuntimeError('Set ATHENA_TEST_SELF_HOST_URL to this execution host URL before running the network isolation smoke test.')
    require_free_threading()
    config = Config.from_env()
    if not (config.data_dir / ".athena-execution-fixture").is_file():
        raise SystemExit("This smoke must use a disposable, explicitly marked data directory.")
    config.prepare()
    store = Store(config.data_dir / "smoke.sqlite3")
    await store.initialize()
    await initialize_files(store)
    runner = ToolRunner(config, store)
    keep_browser = False
    try:
        if '--browser-fixture' in sys.argv:
            fixture = json.loads((config.data_dir / 'browser-fixture.json').read_text())
            page = None
            async def operate(action, **fields):
                nonlocal page
                args = {'action': action, **fields}
                if page and action != 'open':
                    args.update(tab_id=page['tab_id'], observation_id=page['observation_id'])
                page = await runner.browser.operate(args, fixture['session'], None)
                assert page.get('ok'), page
                return page
            def ref(name):
                line = next(line for line in page['text'].splitlines() if f'"{name}"' in line and '[ref=' in line)
                return '@' + re.search(r'\[ref=(e\d+)\]', line).group(1)
            await operate('open', url='http://127.0.0.1:8123/')
            assert 'Guest' in page['text']
            await operate('fill', ref=ref('Name'), text='Athena Unicode 雅典娜')
            await operate('click', ref=ref('Submit'))
            for _ in range(15):
                if 'Submitted once' in page['text']:
                    break
                await asyncio.sleep(.05)
                await operate('snapshot')
            assert 'Submitted once' in page['text']
            download = await operate('download', ref=ref('Download statement'), name='statement.csv')
            record = (await store.rows('SELECT * FROM file_entries WHERE id=?', (download['file']['id'],)))[0]
            handle, _ = _open_file(config, record)
            with handle:
                assert handle.read() == b'item,amount\nfixture,42\n'
            image = await operate('screenshot')
            record = (await store.rows('SELECT * FROM file_entries WHERE id=?', (image['file']['id'],)))[0]
            handle, _ = _open_file(config, record)
            with handle:
                assert handle.read(8) == b'\x89PNG\r\n\x1a\n'
            await runner.browser.driver.close(fixture['owner'])
            for _ in range(30):
                if not await runner.browser.driver.identity(fixture['owner']):
                    break
                await asyncio.sleep(.05)
            page = None
            await operate('open', url='http://127.0.0.1:8123/')
            assert 'Remembered account' in page['text']
            await operate('open', url='http://127.0.0.1:8123/receipts')
            assert '\\"count\\": 1' in page['text'] or '"count": 1' in page['text'], page
            print(json.dumps({'real_form': True, 'post_receipts': 1, 'exact_download': True,
                              'png_screenshot': True, 'cookie_after_browser_restart': True}))
            return
        auth = Auth(store)
        suffix = uuid.uuid4().hex[:8]
        users = [(await auth.register("fixture-" + name + suffix, "fixture-only-password"))[0] for name in ("alice", "bob")]
        sessions = [await store.session(user_id=user['id']) for user in users]
        a, b = sessions
        if '--prepare-browser-fixture' in sys.argv:
            await runner.browser.driver.start(users[0]['id'])
            fixture = {'owner': users[0]['id'], 'session': a}
            (config.data_dir / 'browser-fixture.json').write_text(json.dumps(fixture))
            keep_browser = True
            print(json.dumps(fixture))
            return
        source = '''import os,sys,sysconfig,pathlib
from openpyxl import Workbook,load_workbook
assert sysconfig.get_config_var('Py_GIL_DISABLED') and not sys._is_gil_enabled()
assert 'DEEPSEEK_API_KEY' not in os.environ
assert not pathlib.Path('/run/athena-execution/broker.sock').exists()
assert not pathlib.Path('/var/run/docker.sock').exists()
assert not pathlib.Path('/opt/athena/backend').exists()
assert not pathlib.Path('/data/smoke.sqlite3').exists()
w=Workbook();w.active.append(['Item','Value']);w.active.append(['Verified',42]);w.save('result.xlsx')
assert load_workbook('result.xlsx').active['B2'].value==42
pathlib.Path('private-note.txt').write_text('account-isolation-fixture')
print('verified-xlsx-and-boundaries')
'''
        output = await runner.execute('exec', {'command': 'python -c ' + shlex.quote(source)}, a, None)
        assert output['exit_code'] == 0, output
        assert 'verified-xlsx' in output['output']
        await discover_files(config, store, users[0]['id'], a)
        assert (config.data_dir / 'files' / users[0]['id'] / a / 'result.xlsx').is_file()
        other_file = str(config.data_dir / 'files' / users[0]['id'] / a / 'private-note.txt')
        check = 'from pathlib import Path; assert not Path(' + repr(other_file) + ').exists(); print("owner-isolated")'
        other = await runner.execute('exec', {'command': 'python -c ' + shlex.quote(check)}, b, None)
        assert other['exit_code'] == 0 and 'owner-isolated' in other['output'], other
        public = await runner.execute('exec', {'command': 'curl --fail --max-time 20 --silent https://example.com'}, a, None)
        assert public['exit_code'] == 0 and 'Example Domain' in public['output'], public
        private = await runner.execute('exec', {'command': 'curl --fail --max-time 2 --silent http://169.254.169.254/'}, a, None)
        assert private['exit_code'] != 0, private
        self_host = await runner.execute('exec', {'command': 'curl --fail --max-time 2 --silent ' + shlex.quote(self_host_url)}, a, None)
        assert self_host['exit_code'] != 0, self_host
        peer = os.environ.get('ATHENA_TEST_PEER_URL')
        if peer:
            neighbor = await runner.execute('exec', {'command': 'curl --fail --max-time 2 --silent ' + shlex.quote(peer)}, a, None)
            assert neighbor['exit_code'] != 0, neighbor
        large = await runner.execute('exec', {'command': "python -c 'print(\"evidence\" * 10000)'"}, a, None)
        original = await runner.execute('runtime', {'op': 'read_artifact', 'args': {'artifact_id': large['artifact_id'], 'offset': 65000}}, a, None)
        assert large['exit_code'] == 0 and len(large['output']) < 13000 and original['total_chars'] == 80001
        timed = await runner.execute('exec', {'command': 'printf started; sleep 30', 'timeout_seconds': 1}, a, None)
        assert timed['error'] and 'started' in timed['output'], timed
        heartbeat = config.data_dir / 'files' / users[0]['id'] / a / 'heartbeat.txt'
        child = "import pathlib,time\np=pathlib.Path('heartbeat.txt')\nwhile True:\n p.write_text(str(time.time_ns()));time.sleep(.05)"
        detached = "import subprocess,sys,time;subprocess.Popen([sys.executable,'-c'," + repr(child) + "],start_new_session=True);time.sleep(30)"
        active = asyncio.create_task(runner.execute('exec', {'command': 'python -c ' + shlex.quote(detached)}, a, None))
        for _ in range(100):
            if heartbeat.exists():
                break
            await asyncio.sleep(.05)
        assert heartbeat.exists(), 'Detached test child never started.'
        active.cancel()
        await asyncio.gather(active, return_exceptions=True)
        await asyncio.sleep(.5)
        stamp = heartbeat.stat().st_mtime_ns
        await asyncio.sleep(1)
        assert heartbeat.stat().st_mtime_ns == stamp, 'A detached child survived Stop.'
        page = await runner.execute('runtime', {'op': 'browser', 'args': {'action': 'open', 'url': 'https://example.com'}}, a, None)
        assert page['ok'] and len(page['text']) > 100 and page['url'].startswith('https://example.com'), page
        picture = await runner.execute('runtime', {'op': 'browser', 'args': {'action': 'screenshot', 'tab_id': page['tab_id']}}, a, None)
        assert picture['ok'], picture
        try:
            await runner.execute('runtime', {'op': 'browser', 'args': {'action': 'snapshot', 'tab_id': page['tab_id']}}, b, None)
        except ValueError:
            pass
        else:
            raise AssertionError('A browser tab crossed the account boundary.')
        print(json.dumps({'commands': True, 'xlsx_roundtrip': True, 'file_isolation': True,
                          'public_http': True, 'metadata_blocked': True, 'recoverable_output': True,
                          'timeout': True, 'detached_child_cancelled': True, 'host_services_blocked': True,
                          'same_bridge_blocked': bool(peer),
                          'browser_page': True, 'screenshot': True, 'tab_isolation': True}))
    finally:
        if not keep_browser:
            await runner.browser.close()
        await store.close()


asyncio.run(main())
